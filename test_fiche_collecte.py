"""Tests : la fiche de collecte générée est acceptée par l'import (tools/build_fiche.py).
Lancer : python -m unittest test_fiche_collecte -v   (nécessite openpyxl : requirements-dev.txt)"""
import csv
import datetime
import io
import os
import sys
import tempfile
import unittest
import unittest.mock

try:
    import openpyxl
except ImportError:  # outil de développement : absent en production
    openpyxl = None

import import_prices
import pricing

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tools'))


@unittest.skipIf(openpyxl is None, 'openpyxl non installé')
class TestFicheCollecte(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import build_fiche
        cls.build = build_fiche
        cls.tmp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.tmp.cleanup)
        cls.stores = ['Carrefour', 'Cap Sud', 'Casino']
        cls.rows = build_fiche.build_rows(cls.stores)
        build_fiche.write_csv(os.path.join(cls.tmp.name, 'f.csv'), cls.rows)
        build_fiche.write_workbook(os.path.join(cls.tmp.name, 'f.xlsx'), cls.rows, cls.stores)

    def workbook(self):
        return openpyxl.load_workbook(os.path.join(self.tmp.name, 'f.xlsx'))

    def test_blank_sheet_imports_as_all_skipped_not_as_errors(self):
        with open(os.path.join(self.tmp.name, 'f.csv'), 'rb') as f:
            rows, errors, skipped = import_prices.parse_csv(io.StringIO(import_prices.decode_csv(f.read())))
        self.assertEqual((rows, errors, skipped), ([], [], len(self.rows)))

    def test_saisie_columns_match_the_importer_and_the_csv(self):
        header = [c.value for c in self.workbook()['Saisie'][1]]
        self.assertEqual(header, self.build.COLUMNS)
        for required in import_prices.REQUIRED:
            self.assertIn(required, header)

    def test_sheets_and_active_tab(self):
        wb = self.workbook()
        self.assertEqual(wb.sheetnames, ['Lisez-moi', 'Saisie', 'Suivi'])
        self.assertEqual(wb.active.title, 'Saisie')  # seul l'onglet actif part en CSV

    def test_every_product_gets_a_comparable_unit_price_from_its_name_and_unit(self):
        for name, unit in self.build.PRODUCTS:
            price = pricing.unit_price({'article': name, 'prix': 1000, 'unite': unit, 'source': 'manuel'})
            self.assertIsNotNone(price[0], f'{name} : prix unitaire impossible')

    def test_every_row_uses_valid_static_values(self):
        for name, store, price, unit, day, source, status, *_ in self.rows:
            self.assertIn(unit, pricing.UNITES, name)
            self.assertIn(source, pricing.SOURCES - {'exemple'})
            self.assertIn(status, pricing.STATUTS - {'donnee_exemple'})
        self.assertEqual(len({(r[0], r[1]) for r in self.rows}), len(self.rows), 'doublon produit/magasin')

    def test_instructions_do_not_tell_collectors_to_write_text_in_url(self):
        text = ' '.join(str(c.value) for row in self.workbook()['Lisez-moi'].iter_rows() for c in row if c.value)
        self.assertNotIn('colonne url', text)  # une url non http(s) ferait rejeter tout l'import

    def test_french_excel_export_of_a_filled_sheet_is_imported(self):
        """Remplissage réaliste : prix avec espace et « FCFA », date JJ/MM/AAAA, séparateur « ; », note et lieu."""
        today = datetime.date.today()
        out = io.StringIO()
        writer = csv.writer(out, delimiter=';')
        writer.writerow(self.build.COLUMNS)
        filled = [list(r) for r in self.rows]
        filled[0][2:5] = ['2 500 FCFA', 'kg', today.strftime('%d/%m/%Y')]
        filled[0][9:] = ['Carrefour Cocody · Abidjan', 'Uncle Sam, promo']
        filled[1][2:5] = ['2650', 'kg', today.isoformat()]
        writer.writerows(filled)
        rows, errors, skipped = import_prices.parse_csv(io.StringIO(out.getvalue()))
        self.assertEqual(errors, [])
        self.assertEqual([r['prix'] for r in rows], [2500, 2650])
        self.assertEqual(rows[0]['date_releve'], today.isoformat())
        self.assertEqual(rows[0]['lieu'], 'Carrefour Cocody · Abidjan')
        self.assertEqual(skipped, len(self.rows) - 2)

    def test_rejects_duplicate_or_empty_store_lists(self):
        import build_fiche
        with unittest.mock.patch.object(sys, 'argv', ['build_fiche', '--stores', 'A,A', '--out', self.tmp.name]):
            with self.assertRaises(SystemExit):
                build_fiche.main()


if __name__ == '__main__':
    unittest.main()
