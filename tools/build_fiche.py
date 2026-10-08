"""Génère la fiche de collecte des prix : data/fiche_collecte.xlsx (à remplir) et data/fiche_collecte.csv (vierge).

    python tools/build_fiche.py                              # 3 magasins, 30 produits
    python tools/build_fiche.py --stores "Carrefour,Casino"  # autres magasins
    python tools/build_fiche.py --out /tmp/fiche             # autre dossier de sortie

Le classeur contient trois feuilles :
  - « Lisez-moi » : règles de relevé et marche à suivre jusqu'à l'import ;
  - « Saisie »    : une ligne par produit et par magasin, avec exactement les colonnes attendues par
                    `import_prices.py` et par /admin (Enregistrer sous → CSV) ;
  - « Suivi »     : avancement par magasin (formules, mises à jour pendant la saisie).

Dépendance : openpyxl (outil de développement uniquement, voir requirements-dev.txt).
"""
import argparse
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from pricing import MAX_PRICE_FCFA, UNITES  # noqa: E402  (mêmes bornes que l'import)

# 9 premières colonnes : le format historique de l'import. `lieu` est conservé et affiché aux utilisateurs ;
# `note` est réservée au collecteur et à l'administrateur (l'import l'ignore).
COLUMNS = ['article', 'supermarche', 'prix', 'unite', 'date_releve', 'source', 'statut', 'url', 'image_url', 'lieu', 'note']
DEFAULT_STORES = ['Carrefour', 'Cap Sud', 'Casino']

# (nom avec le FORMAT, unité). Le prix attendu est celui du format écrit dans le nom (voir « Lisez-moi »).
PRODUCTS = [
    ('Riz parfumé 5kg', 'kg'), ('Riz brisé 5kg', 'kg'),
    ('Huile végétale 1L', 'L'), ('Huile végétale 5L', 'L'),
    ('Sucre en poudre 1kg', 'kg'), ('Farine de blé 1kg', 'kg'), ('Spaghetti 500g', 'kg'), ('Sel fin 1kg', 'kg'),
    ('Lait en poudre 400g', 'kg'), ('Lait concentré sucré 397g', 'kg'), ('Lait UHT demi-écrémé 1L', 'L'),
    ('Concentré de tomate 400g', 'kg'), ("Sardines à l'huile 125g", 'kg'), ('Maquereau en boîte 155g', 'kg'),
    ('Oeufs plaquette de 30', 'unité'), ('Pain baguette', 'unité'), ('Beurre 250g', 'kg'), ('Margarine 500g', 'kg'),
    ('Café soluble 200g', 'kg'), ('Eau minérale 1,5L', 'L'), ("Jus d'orange 1L", 'L'),
    ('Poulet entier congelé 1kg', 'kg'), ('Oignons 1kg', 'kg'), ('Tomates fraîches 1kg', 'kg'),
    ('Pommes de terre 1kg', 'kg'), ('Igname 1kg', 'kg'), ('Banane plantain 1kg', 'kg'), ('Attiéké 500g', 'kg'),
    ('Cubes de bouillon (boîte de 100)', 'unité'), ('Savon de ménage (pain)', 'unité'),
]

RULES = [
    ('À EMPORTER', None),
    ('•', "Ce classeur (sur téléphone ou ordinateur), ou son impression ; un stylo ; l'appli Appareil photo (facultatif : photo de l'étiquette)."),
    ('RÈGLES DE RELEVÉ (elles évitent 90 % des erreurs)', None),
    ('1.', "Notez le prix de l'ÉTIQUETTE EN RAYON, pour le FORMAT écrit dans la colonne « article » : « Riz parfumé 5kg » = prix du sac de 5 kg, PAS le prix au kilo. L'application calcule elle-même le prix au kilo."),
    ('2.', "NE MODIFIEZ PAS la colonne « article » : la comparaison entre magasins se fait sur le nom exact. Prenez dans chaque magasin la marque la plus vendue de ce produit, et notez-la dans la colonne « note » (ex. « Uncle Sam »). Si une seule marque très différente est disponible, notez-la aussi : l'administrateur décidera."),
    ('3.', "Prix PROMOTIONNEL : notez le prix payé en caisse aujourd'hui (le prix barré n'est pas le prix) et écrivez « promo » dans la colonne « note »."),
    ('4.', "Produit absent ou en rupture : LAISSEZ LA CASE « prix » VIDE. Les lignes sans prix sont ignorées à l'import, jamais refusées."),
    ('5.', "Prix au détail seulement. Ne notez pas un prix « de gros » ni un prix remisé par carte de fidélité."),
    ('5 bis.', "Colonne « lieu » : le point de vente précis et sa commune, par exemple « Carrefour Cocody · Abidjan » (copiez-collez la même valeur sur toutes les lignes de ce magasin). Les utilisateurs la voient : elle les aide à savoir où aller."),
    ('6.', "Date : le jour où VOUS avez vu le prix (jamais une date future). Format conseillé : AAAA-MM-JJ, par exemple 2026-10-08. Le format JJ/MM/AAAA est aussi accepté."),
    ('7.', "Prix : un nombre entier en FCFA, sans texte (2500). « 2 500 » et « 2500 FCFA » sont acceptés ; « 2500,5 » est refusé (aucun arrondi)."),
    ('8.', "Les colonnes source (manuel) et statut (a_verifier) sont déjà remplies : ne les changez pas, sauf avec l'accord de l'administrateur."),
    ('MARCHE À SUIVRE', None),
    ('A.', "Onglet « Saisie » : filtrez la colonne supermarche sur le magasin où vous êtes, puis remplissez prix et date_releve. L'onglet « Suivi » montre l'avancement."),
    ('B.', "Enregistrer sous → type « CSV UTF-8 (délimité par des virgules) » (ou « CSV (séparateur : point-virgule) » sur Excel français : les deux sont acceptés). Seul l'onglet actif « Saisie » est enregistré : vérifiez qu'il est affiché."),
    ('C.', "Administrateur : ouvrez /admin → Importer un CSV. Une simulation s'exécute d'abord et rien n'est écrit si une seule ligne est invalide. En ligne de commande : python import_prices.py releve.csv, puis --apply."),
    ('D.', "Les prix importés sont « à vérifier » : validez-les dans /admin (bouton Valider) quand ils ont été contrôlés."),
    ('QUALITÉ', None),
    ('•', "Objectif de départ : 20 produits × 3 magasins = 60 prix, relevés le même jour (± 2 jours) pour que la comparaison soit juste."),
    ('•', "Un prix qui s'écarte de plus de 20 % de la médiane des autres magasins sera signalé comme anormal dans l'application : revérifiez l'étiquette avant d'envoyer."),
    ('•', "Ne notez jamais de données personnelles (visage, ticket avec carte bancaire ou numéro de client) sur une photo."),
]


def build_rows(stores):
    # Produit par produit : on comparera les magasins côte à côte pendant la saisie.
    return [[name, store, '', unit, '', 'manuel', 'a_verifier', '', '', '', ''] for name, unit in PRODUCTS for store in stores]


def write_csv(path, rows):
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:  # BOM : Excel ouvre les accents correctement
        writer = csv.writer(f)
        writer.writerow(COLUMNS)
        writer.writerows(rows)


def write_workbook(path, rows, stores):
    wb = Workbook()
    green, grey = PatternFill('solid', fgColor='176B45'), PatternFill('solid', fgColor='EEF3EE')
    thin = Side(style='thin', color='C8D4CB')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    guide = wb.active
    guide.title = 'Lisez-moi'
    guide.column_dimensions['A'].width = 6
    guide.column_dimensions['B'].width = 120
    guide['A1'] = 'ComparePrix : fiche de collecte des prix en magasin'
    guide['A1'].font = Font(size=16, bold=True, color='176B45')
    row = 3
    for left, text in RULES:
        if text is None:
            guide.cell(row=row, column=1, value=left).font = Font(bold=True, color='FFFFFF')
            for col in (1, 2):
                guide.cell(row=row, column=col).fill = green
        else:
            guide.cell(row=row, column=1, value=left).alignment = Alignment(vertical='top')
            cell = guide.cell(row=row, column=2, value=text)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            guide.row_dimensions[row].height = 15 * (1 + len(text) // 115)
        row += 1

    sheet = wb.create_sheet('Saisie')
    sheet.append(COLUMNS)
    for line in rows:
        sheet.append(line)
    last = len(rows) + 1
    for col, width in zip('ABCDEFGHIJK', (38, 16, 12, 9, 14, 10, 12, 18, 14, 30, 28)):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True, color='FFFFFF'), green, Alignment(horizontal='center')
    for r in range(2, last + 1):
        for c in range(1, 12):
            sheet.cell(row=r, column=c).border = box
        for c in (1, 2, 6, 7):  # colonnes déjà remplies : grisées pour ne pas les modifier par erreur
            sheet.cell(row=r, column=c).fill = grey
        sheet.cell(row=r, column=3).number_format = '0'
        sheet.cell(row=r, column=5).number_format = 'yyyy-mm-dd'
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = f'A1:K{last}'

    def validation(kind, cells, **kw):
        dv = DataValidation(type=kind, allow_blank=True, showErrorMessage=True, errorStyle='stop', **kw)
        sheet.add_data_validation(dv)
        dv.add(cells)

    validation('whole', f'C2:C{last}', operator='between', formula1='1', formula2=str(MAX_PRICE_FCFA),
               errorTitle='Prix invalide', error='Entrez un nombre entier en FCFA, sans texte (ex. 2500). Laissez vide si le produit est absent.')
    validation('list', f'D2:D{last}', formula1='"' + ','.join(sorted(UNITES)) + '"',
               errorTitle='Unité invalide', error='Choisissez une unité de la liste.')
    validation('date', f'E2:E{last}', operator='greaterThan', formula1='45292',  # 2024-01-01
               errorTitle='Date invalide', error='Entrez la date du relevé, par exemple 2026-10-08 (pas de date future).')
    validation('list', f'F2:F{last}', formula1='"manuel,ticket"', errorTitle='Source invalide', error='manuel ou ticket.')
    validation('list', f'G2:G{last}', formula1='"a_verifier,valide"', errorTitle='Statut invalide', error='a_verifier ou valide.')

    # Retour visuel pendant la saisie : vert = complet, rouge = prix sans date.
    sheet.conditional_formatting.add(f'A2:K{last}', FormulaRule(formula=['AND($C2<>"",$E2="")'], fill=PatternFill('solid', bgColor='F8D7D7')))
    sheet.conditional_formatting.add(f'A2:K{last}', FormulaRule(formula=['AND($C2<>"",$E2<>"")'], fill=PatternFill('solid', bgColor='D9EFDF')))

    follow = wb.create_sheet('Suivi')
    follow.append(['Magasin', 'Lignes à remplir', 'Prix saisis', 'Avancement', 'Prix sans date'])
    for cell in follow[1]:
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), green
    for i, store in enumerate(stores, start=2):
        follow.append([store,
                       f'=COUNTIF(Saisie!$B$2:$B${last},A{i})',
                       f'=COUNTIFS(Saisie!$B$2:$B${last},A{i},Saisie!$C$2:$C${last},"<>")',
                       f'=IF(B{i}=0,0,C{i}/B{i})',
                       f'=COUNTIFS(Saisie!$B$2:$B${last},A{i},Saisie!$C$2:$C${last},"<>",Saisie!$E$2:$E${last},"")'])
        follow.cell(row=i, column=4).number_format = '0%'
    total = len(stores) + 2
    follow.append(['Total', f'=SUM(B2:B{total - 1})', f'=SUM(C2:C{total - 1})', f'=IF(B{total}=0,0,C{total}/B{total})', f'=SUM(E2:E{total - 1})'])
    follow.cell(row=total, column=4).number_format = '0%'
    for cell in follow[total]:
        cell.font = Font(bold=True)
    for col, width in zip('ABCDE', (22, 18, 14, 14, 16)):
        follow.column_dimensions[col].width = width

    wb.active = 1  # s'ouvre sur « Saisie » : c'est l'onglet exporté en CSV
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--stores', default=','.join(DEFAULT_STORES), help='magasins séparés par des virgules')
    parser.add_argument('--out', default=os.path.join(ROOT, 'data'), help='dossier de sortie (data/ par défaut)')
    args = parser.parse_args()
    stores = [s.strip() for s in args.stores.split(',') if s.strip()]
    if not stores or len(set(stores)) != len(stores):
        raise SystemExit('Indiquez au moins un magasin, sans doublon.')
    os.makedirs(args.out, exist_ok=True)
    rows = build_rows(stores)
    write_csv(os.path.join(args.out, 'fiche_collecte.csv'), rows)
    write_workbook(os.path.join(args.out, 'fiche_collecte.xlsx'), rows, stores)
    print(f'{len(rows)} lignes ({len(PRODUCTS)} produits × {len(stores)} magasins) → {args.out}')


if __name__ == '__main__':
    main()
